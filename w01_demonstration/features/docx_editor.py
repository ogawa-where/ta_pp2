#!/usr/bin/env python3
#
# docx_editor.py
# 
# [概要]
# python-docxモジュールを使い，
# 既存のドキュメントファイルを編集するための
# スクリプト．
# 大学の施設利用申請書を編集する機能を実装している．
# 

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from features.copy_template import copy_template

from pathlib import Path
from logging import getLogger
import os

logger = getLogger(__name__)


class DocxEditor:
    def __init__(self, input_docx, output_docx):
        self.output_path = output_docx
        copy_template(input_docx, output_docx)
        self.docx = Document(output_docx)
        
    def read_pargraphs(self):
        '''
        [概要]
        書式内の全ての段落から文字列を抽出してリストに格納する関数．
        なお．この処理ではテーブル(表)のセル情報まで掴めない．

        Return:
            texts (list): 書式内の全ての段落に含まれる文字列情報が入ったリスト
        '''
        try:
            texts = [
                paragraph.text for paragraph in self.docx.paragraphs
            ]
            logger.info("全ての段落からテキストを抽出しました")
            return texts

        except AttributeError:
            logger.error(
                "docxオブジェクトが正しく初期化されていません．"
            )
            raise

        except Exception as e:
            logger.exception(
                f"テキストの抽出中に予期せぬエラーが発生しました: {e}"
            )
            raise
            
    def read_tables(self, table_num=0):
        '''
        [概要]
        書式内のテーブル(表)オブジェクトを取得してリストに格納する．
        そして，任意のテーブル番号(デフォルトは0番目)のオブジェクトを返す関数．

        Args:
            table_num (int): 取得したいテーブル番号を指示(デフォルトは0番目)．

        Return:
            tables[table_num] (obj): 指示された番号のテーブルオブジェクト
        '''
        try:
            tables = self.docx.tables
            logger.info(
                f"{table_num}番目のテーブルオブジェクトを取得しました"
            )
            return tables[table_num]

        except IndexError:
            table_count = len(self.docx.tables)
            logger.error(
                f"指定されたテーブル番号{table_num}は範囲外です．\n"
                f"このドキュメントには{table_count}個のテーブルしかありません"
            )
            raise

        except AttributeError:
            logger.error(
                f"docxオブジェクトが正しく初期化されていません"
            )
            raise

        except Exception as e:
            logger.exception(
                f"テーブルの読み込み中に予期せぬエラーが発生しました: {e}"
            )
            raise

    def read_cells(self, table_num=0):
        '''
        [概要]
        テーブルオブジェクトを生成して，テーブルに含まれる全てのセルの
        テキストを取得して，2次元リストとして返す関数．

        Args: 
            table_num: read_tables()を動かすための引数．

        Return:
            texts (list[list]): 各セルのテキストを格納した2次元リスト．
        '''
        try:
            table = self.read_tables(table_num)
            
            texts = [
                [cell.text for cell in row.cells]
                for row in table.rows
            ]

            logger.info(
                f"テーブル{table_num}から全てのセルのテキストを抽出しました"
            )
            return texts

        except Exception as e:
            logger.exception(
                f"テーブル{table_num}からのセルテキスト抽出に失敗しました"
            )
            raise

    def read_cell_paragraphs(self, table_num=0, row_idx=0, column_idx=0):
        '''
        [概要]
        テーブルオブジェクトを生成して，テーブル内に存在するセルを一つ指定し，
        単一セル内の段落オブジェクトをリストにまとめて返す関数．

        Args:
            table_num: read_tables()を動かすための引数．
            row_idx: セル指定における行を示す．
            column_idx: セル指定における列を示す．

        Return:
            table().cell().paragraphs(): 段落オブジェクトが格納されたリスト．
        '''
        return self.read_tables(table_num).cell(row_idx, column_idx).paragraphs
    
    def replace_paragraph_text(self, paragraph, placeholders):
        '''
        [概要]
        段落オブジェクトと置換情報を利用して，書式内の狙った箇所だけを
        編集する関数．
        段落情報を対象とするため，セル内の段落情報を取得できていれば
        テーブル内のセル群にも対応できる．
        
        Args:
            paragraph (obj): 置換対象となる段落オブジェクト．
            placeholder (dict): 置換対象と置換後のテキストデータを格納した辞書
        '''
        try:
            for key, value in placeholders.items():
                placeholder = "{" + key + "}"
                for run in paragraph.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, str(value))

            logger.debug("段落内のテキスト置換処理が正常に完了しました")

        except AttributeError:
            logger.error(
                "不正なオブジェクトが入力されました．\n"
                "引数 'paragraph' が段落オブジェクトであるか，"
                "'placeholders' が辞書型であるか確認してください．"
            )
            raise

        except TypeError:
            logger.error(
                "引数 'placeholders' のデータ型が不適切です．\n"
                "キーと値を持つ辞書を指定してください．"
            )
            raise

        except Exception as e:
            logger.exception(
                f"テキストの置換中に予期せぬエラーが発生しました: {e}"
            )
            raise

    def make_facilty_usage_form(self, placeholders):
        '''
        [概要]
        秋田県立大学本荘キャンパスにて，体育館やトレーニングルームといった
        施設を利用する際に必要な申請書の書式を作成して保存するための関数．

        Args:
            placeholders: replace_paragraph_text()を動かすために使用する引数．
        '''
        try:
            paragraphs = self.read_cell_paragraphs()

            for paragraph in paragraphs:
                self.replace_paragraph_text(paragraph, placeholders)

            self.docx.save(self.output_path)
            logger.info("ドキュメントの複製・編集・保存が完了しました")

            return True

        except Exception as e:
            logger.error(
                f"ドキュメントの操作中に予期せぬエラーが発生しました: {e}"
            )
            return None

    
if __name__ == "__main__":
    "--- 手動テスト ---"
    from setup_logging import setup_logging

    setup_logging("../.config/logging_config.yml")

    logger.debug("=== DocxEditor 手動テスト開始 ===")
    
    os.makedirs(
        "output", exist_ok=True
    )

    placeholders = {
        "wareki": "7",
        "mm": "9",
        "DD": "23",
        "department": "経営システム工学専攻",
        "student_number": "M26D003",
        "student_name": "吉田 快",
        "group_name": "  ",
        "teacher_name": "  ",
        "A_or_P": "午後",
        "sH": "18",
        "sM": "30",
        "eH": "20",
        "eM": "00",
        "purpose_of_use": "筋トレ",
        "total_num": "1"
    }
        
    input_docx = "../input/facility_use_form.docx"
    output_docx = "./output/output_form.docx"
    editor = DocxEditor(input_docx, output_docx)
    editor.make_facilty_usage_form(placeholders)

    logger.debug("=== DocxEditor 手動テスト終了 ===")
